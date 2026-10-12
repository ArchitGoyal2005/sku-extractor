
import io
import logging
import os
import re
import secrets
import threading
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

import pymupdf
from flask import (
    Flask,
    Request as FlaskRequest,
    abort,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from werkzeug.utils import secure_filename


# -------------------------------------------------------------------
# Configuration
# -------------------------------------------------------------------

MAX_MB = int(os.environ.get("MAX_UPLOAD_MB", "50"))
MAX_BYTES = MAX_MB * 1024 * 1024

JOB_TTL_SECONDS = int(
    os.environ.get("JOB_TTL_SECONDS", "1800")
)

MAX_IN_MEMORY_JOBS = int(
    os.environ.get("MAX_IN_MEMORY_JOBS", "20")
)

MAX_PAGES = int(os.environ.get("MAX_PAGES", "2000"))


# -------------------------------------------------------------------
# Logging
# -------------------------------------------------------------------

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
)

logger = logging.getLogger("sku-extractor")


# -------------------------------------------------------------------
# Memory-backed upload handling
# -------------------------------------------------------------------

class MemoryRequest(FlaskRequest):
    """
    Store uploaded multipart file data in memory instead of using
    Werkzeug's default temporary-file upload stream.
    """

    def _get_file_stream(
        self,
        total_content_length,
        content_type,
        filename=None,
        content_length=None,
    ):
        return io.BytesIO()


app = Flask(__name__)
app.request_class = MemoryRequest

app.secret_key = os.environ.get("FLASK_SECRET_KEY")

if not app.secret_key:
    raise RuntimeError(
        "Set the FLASK_SECRET_KEY environment variable before starting the app."
    )

app.config["MAX_CONTENT_LENGTH"] = MAX_BYTES
app.config["MAX_FORM_MEMORY_SIZE"] = MAX_BYTES
app.config["MAX_FORM_PARTS"] = 10


# -------------------------------------------------------------------
# In-memory job cache
# -------------------------------------------------------------------

# No application-managed PDF or Excel files are intentionally saved
# to disk. Results and generated Excel bytes remain in process memory.
JOBS = {}
JOBS_LOCK = threading.Lock()


def cleanup_expired_jobs():
    """Remove expired jobs and enforce the maximum cache size."""

    cutoff = time.time() - JOB_TTL_SECONDS

    with JOBS_LOCK:
        expired = [
            job_id
            for job_id, job in JOBS.items()
            if job["created"] < cutoff
        ]

        for job_id in expired:
            JOBS.pop(job_id, None)

        while len(JOBS) > MAX_IN_MEMORY_JOBS:
            oldest = min(
                JOBS,
                key=lambda key: JOBS[key]["created"],
            )
            JOBS.pop(oldest, None)


# -------------------------------------------------------------------
# SKU extraction
# -------------------------------------------------------------------

def extract_sku(page):
    """
    Extract the SKU from a page's selectable PDF text.

    Assumption:
    The SKU is the first non-empty line after the Product Details
    table header and before TAX INVOICE.

    Adjust this pattern if your PDF uses a different text layout.
    """

    text = page.get_text("text")

    match = re.search(
        r"Product\s+Details\s+SKU\s+Size\s+Qty\s+Color\s+"
        r"Order\s+No\.?\s*(.*?)\s*TAX\s+INVOICE",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )

    if not match:
        return None

    lines = [
        line.strip()
        for line in match.group(1).splitlines()
        if line.strip()
    ]

    if not lines:
        return None

    # Skip a repeated table header if it appears in the captured text.
    while lines and lines[0].casefold() in {
        "sku",
        "size",
        "qty",
        "color",
        "order no",
        "order no.",
    }:
        lines.pop(0)

    if not lines:
        return None

    return lines[0]


# -------------------------------------------------------------------
# Excel generation
# -------------------------------------------------------------------

def create_excel_bytes(counts):
    """Generate the Excel report entirely in memory."""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SKU Report"

    worksheet.append(["SKU ID", "Count"])

    sorted_counts = sorted(
        counts.items(),
        key=lambda item: (
            -item[1],
            item[0].casefold(),
        ),
    )

    for sku, count in sorted_counts:
        worksheet.append([str(sku), int(count)])

        # Preserve SKU values as text in Excel.
        sku_cell = worksheet.cell(
            row=worksheet.max_row,
            column=1,
        )
        sku_cell.number_format = "@"
        sku_cell.data_type = "s"

        count_cell = worksheet.cell(
            row=worksheet.max_row,
            column=2,
        )
        count_cell.number_format = "0"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    worksheet.column_dimensions["A"].width = 34
    worksheet.column_dimensions["B"].width = 12

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    if worksheet.max_row > 1:
        table = Table(
            displayName="SKUReport",
            ref=worksheet.dimensions,
        )

        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showRowStripes=True,
        )

        worksheet.add_table(table)

    output = io.BytesIO()

    workbook.save(output)
    output.seek(0)

    excel_bytes = output.getvalue()
    output.close()

    workbook.close()

    return excel_bytes


# -------------------------------------------------------------------
# PDF processing
# -------------------------------------------------------------------

def process_pdf_bytes(pdf_bytes, original_name):
    """Process PDF bytes and return result data plus Excel bytes."""

    logs = []
    occurrences = []
    missing = []

    def log(message):
        line = (
            f"[{datetime.now().strftime('%H:%M:%S')}] "
            f"{message}"
        )

        logs.append(line)
        logger.info("%s", message)

    log(f"Opening uploaded PDF in memory: {original_name}")

    with pymupdf.open(
        stream=pdf_bytes,
        filetype="pdf",
    ) as pdf:

        total = len(pdf)

        if total == 0:
            raise ValueError("The PDF contains no pages.")

        if total > MAX_PAGES:
            raise ValueError(
                f"The PDF contains {total} pages. "
                f"The maximum allowed is {MAX_PAGES}."
            )

        log(
            f"PDF opened successfully; "
            f"{total} page(s) detected."
        )

        for page_no, page in enumerate(pdf, start=1):
            sku = extract_sku(page)

            if sku is None:
                missing.append({
                    "page": page_no,
                    "reason": "SKU section missing or unreadable",
                })

                log(
                    f"Page {page_no}: SKU not found; "
                    "flagged for review."
                )

            else:
                occurrences.append({
                    "sku": sku,
                    "page": page_no,
                })

                log(
                    f"Page {page_no}: "
                    f"extracted SKU '{sku}'."
                )

    counts = Counter(
        item["sku"]
        for item in occurrences
    )

    excel_bytes = create_excel_bytes(counts)

    log(
        f"Finished: {len(occurrences)} SKU occurrence(s), "
        f"{len(counts)} unique SKU(s)."
    )

    log(f"Pages requiring review: {len(missing)}.")

    log(
        "Excel report generated in memory; sorted by "
        "Count descending, then SKU ID alphabetically."
    )

    result = {
        "original_name": original_name,
        "total_pages": total,
        "occurrences": len(occurrences),
        "unique_count": len(counts),
        "missing_count": len(missing),
        "missing_pages": missing,
        "counts": sorted(
            counts.items(),
            key=lambda item: (
                -item[1],
                item[0].casefold(),
            ),
        ),
        "logs": logs,
        "created_label": datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        ),
    }

    return result, excel_bytes


# -------------------------------------------------------------------
# Error handling
# -------------------------------------------------------------------

@app.errorhandler(413)
def too_large(_error):
    flash(
        f"File too large. Maximum upload size is {MAX_MB} MB.",
        "error",
    )
    return redirect(url_for("index"))


# -------------------------------------------------------------------
# Home page
# -------------------------------------------------------------------

@app.get("/")
def index():
    cleanup_expired_jobs()

    return render_template(
        "index.html",
        max_mb=MAX_MB,
    )


# -------------------------------------------------------------------
# Upload and process
# -------------------------------------------------------------------

@app.post("/upload")
def upload():
    cleanup_expired_jobs()

    uploaded = request.files.get("pdf")

    if not uploaded or not uploaded.filename:
        flash("Choose a PDF file first.", "error")
        return redirect(url_for("index"))

    original_name = secure_filename(
        uploaded.filename
    ) or "uploaded.pdf"

    if Path(original_name).suffix.lower() != ".pdf":
        flash("Only PDF files are supported.", "error")
        return redirect(url_for("index"))

    pdf_bytes = None

    try:
        # Read the uploaded PDF into memory.
        pdf_bytes = uploaded.stream.read(MAX_BYTES + 1)

        if len(pdf_bytes) > MAX_BYTES:
            flash(
                f"File too large. Maximum upload size is {MAX_MB} MB.",
                "error",
            )
            return redirect(url_for("index"))

        # Reject empty uploads early.
        if not pdf_bytes:
            flash("The uploaded PDF is empty.", "error")
            return redirect(url_for("index"))

        result, excel_bytes = process_pdf_bytes(
            pdf_bytes,
            original_name,
        )

        # Drop the application's reference to the raw PDF bytes.
        del pdf_bytes
        pdf_bytes = None

        # Establish an owner token for this browser session.
        if "owner_token" not in session:
            session["owner_token"] = secrets.token_urlsafe(24)

        owner_token = session["owner_token"]
        job_id = secrets.token_urlsafe(24)

        with JOBS_LOCK:
            JOBS[job_id] = {
                "created": time.time(),
                "owner": owner_token,
                "result": result,
                "excel": excel_bytes,
            }

        session["active_job"] = job_id

        return render_template(
            "result.html",
            result=result,
            job_id=job_id,
        )

    except Exception:
        logger.exception("Failed to process uploaded PDF")

        flash(
            "Could not process that PDF. "
            "Please verify that it is a valid, text-readable PDF.",
            "error",
        )

        return redirect(url_for("index"))

    finally:
        # Release the local reference if an exception occurred.
        pdf_bytes = None
        uploaded.close()


# -------------------------------------------------------------------
# Job access control
# -------------------------------------------------------------------

def get_owned_job(job_id):
    """Return a job only if it belongs to the current browser session."""

    cleanup_expired_jobs()

    owner_token = session.get("owner_token")

    with JOBS_LOCK:
        job = JOBS.get(job_id)

        if not job:
            abort(404)

        if not owner_token or job["owner"] != owner_token:
            abort(404)

        return job


# -------------------------------------------------------------------
# View previous result
# -------------------------------------------------------------------

@app.get("/jobs/<job_id>")
def show_job(job_id):
    job = get_owned_job(job_id)

    session["active_job"] = job_id

    return render_template(
        "result.html",
        result=job["result"],
        job_id=job_id,
    )


# -------------------------------------------------------------------
# Download Excel
# -------------------------------------------------------------------

@app.get("/jobs/<job_id>/download")
def download(job_id):
    job = get_owned_job(job_id)

    # Create a response stream from the Excel bytes in memory.
    stream = io.BytesIO(job["excel"])

    response = send_file(
        stream,
        as_attachment=True,
        download_name="sku_report.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        max_age=0,
    )

    response.headers["Cache-Control"] = "no-store"

    return response


# -------------------------------------------------------------------
# Delete a job
# -------------------------------------------------------------------

@app.post("/jobs/<job_id>/delete")
def delete(job_id):
    get_owned_job(job_id)

    with JOBS_LOCK:
        JOBS.pop(job_id, None)

    if session.get("active_job") == job_id:
        session.pop("active_job", None)

    flash(
        "Job deleted from server memory. "
        "The application does not intentionally save uploaded "
        "PDF or generated Excel files to disk.",
        "success",
    )

    return redirect(url_for("index"))


# -------------------------------------------------------------------
# Clear the current browser session's jobs
# -------------------------------------------------------------------

@app.post("/clear")
def clear():
    owner_token = session.get("owner_token")

    if owner_token:
        with JOBS_LOCK:
            owned_job_ids = [
                job_id
                for job_id, job in JOBS.items()
                if job["owner"] == owner_token
            ]

            for job_id in owned_job_ids:
                JOBS.pop(job_id, None)

    session.pop("active_job", None)

    flash(
        "Your cached jobs have been removed from server memory.",
        "success",
    )

    return redirect(url_for("index"))


# -------------------------------------------------------------------
# Start server
# -------------------------------------------------------------------

if __name__ == "__main__":
    app.run(
        host=os.environ.get("HOST", "127.0.0.1"),
        port=int(os.environ.get("PORT", "5000")),
        debug=False,
    )
